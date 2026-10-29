"""Build the blind relabelling pack for an independent (non-author) expert.

Why: the attribution tests use the uploader's own "pilot" or "drone" rating. An earlier review workbook
(make_review_sheet.py) was never completed and was not blind in practice: it showed the uploader's notes and linked to
Flight Review, which displays the rating. This pack shows only what the log itself records.

SAMPLE: all 36 pilot-rated crashes of the 148 modelled crash logs, plus 36 of the 112 aircraft-rated ones drawn at random
(seed 2080), so that the expert cannot infer the rating from the base rate; shuffled (same seed). Case numbers only.

PER CASE (one PNG, drawn offline from the .ulg with pyulog): over the whole armed period and, enlarged, the last 60 s
before the end of the armed period: roll and pitch with their setpoints; altitude and vertical speed; pilot stick
inputs (manual_control_setpoint); actuator outputs; battery voltage; accelerometer magnitude; flight-mode bands; the
last 15 logged warning/error messages (text only). Also given: vehicle type (index mav_type) and armed duration.
NOT given: log ID, date, vehicle or pilot name, uploader rating, uploader notes, Flight Review link, any rule or model
output, the computed crash moment.

Outputs (folder git-ignored; the key must not be opened by the expert):
  OUT_DIR/cases/case_NN.png, OUT_DIR/relabel_workbook.xlsx (Instructions + Review sheets), OUT_DIR/key_DO_NOT_OPEN.csv

    python research/px4_log_study/make_relabel_pack.py DATA_ROOT OUT_DIR
"""
import csv
import gzip
import json
import os
import random
import sys

import numpy as np

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.append(HERE)
import postcrash_audit as pca

RESULTS = os.path.join(HERE, "..", "results", "px4_log_study")
NAV = {0: "MANUAL", 1: "ALTCTL", 2: "POSCTL", 3: "MISSION", 4: "LOITER", 5: "RTL", 10: "ACRO", 12: "DESCEND",
       13: "TERMINATION", 14: "OFFBOARD", 15: "STABILIZED", 17: "TAKEOFF", 18: "LAND", 19: "FOLLOW", 20: "PRECLAND",
       21: "ORBIT", 22: "VTOL_TAKEOFF"}
VERDICTS = ["pilot", "aircraft", "environment", "can't tell"]
CAUSES = ["pilot input / handling", "wrong mode or mode switch", "flew into obstacle / ground", "battery / power",
          "motor / ESC / propeller", "vibration / mounting", "GPS / compass / estimator", "radio or data link",
          "tuning / configuration", "software / firmware", "airframe / mechanical", "wind / environment",
          "no crash visible", "not enough data"]


def sample(seed=2080):
    crash = {}
    for line in open(os.path.join(RESULTS, "leak_curve_features.jsonl")):
        r = json.loads(line)
        if r["group"] == "crash" and r["margin"] == 0:
            crash[r["log"]] = r["label"]
    pilot = sorted(l for l, v in crash.items() if v == "pilot")
    drone = sorted(l for l, v in crash.items() if v == "drone")
    rng = random.Random(seed)
    s = pilot + rng.sample(drone, len(pilot))
    rng.shuffle(s)
    return [(l, crash[l]) for l in s]


def panel(path, out_png, title):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from pyulog import ULog
    from pilot_vs_drone import armed_period, topic

    u = ULog(path)
    st = topic(u, "vehicle_status")
    t_arm, t_end = armed_period(st)
    T = lambda d: (d["timestamp"] - t_arm) / 1e6
    dur = (t_end - t_arm) / 1e6
    get = lambda name: topic(u, name)
    att, asp = get("vehicle_attitude"), get("vehicle_attitude_setpoint")
    lp, mc, bat = get("vehicle_local_position"), get("manual_control_setpoint"), get("battery_status")
    act = get("actuator_motors") or None
    if act is None:
        act = get("actuator_outputs")
    acc = get("sensor_combined")

    def euler(q0, q1, q2, q3):
        roll = np.degrees(np.arctan2(2 * (q0 * q1 + q2 * q3), 1 - 2 * (q1 ** 2 + q2 ** 2)))
        pitch = np.degrees(np.arcsin(np.clip(2 * (q0 * q2 - q3 * q1), -1, 1)))
        return roll, pitch

    fig, axes = plt.subplots(6, 2, figsize=(15, 15), sharex="col", gridspec_kw={"width_ratios": [1.3, 1]})
    for col, (a, b) in enumerate(((0, dur), (max(0, dur - 60), dur + 2))):
        ax = axes[:, col]
        if att is not None:
            r, p = euler(att["q[0]"], att["q[1]"], att["q[2]"], att["q[3]"])
            ax[0].plot(T(att), r, lw=.7, color="C0", label="roll")
            ax[0].plot(T(att), p, lw=.7, color="C1", label="pitch")
        if asp is not None and "q_d[0]" in asp:
            r, p = euler(asp["q_d[0]"], asp["q_d[1]"], asp["q_d[2]"], asp["q_d[3]"])
            ax[0].plot(T(asp), r, lw=.6, ls="--", color="C0", alpha=.7, label="roll setpoint")
            ax[0].plot(T(asp), p, lw=.6, ls="--", color="C1", alpha=.7, label="pitch setpoint")
        ax[0].set_ylabel("deg")
        if lp is not None:
            ax[1].plot(T(lp), -lp["z"], lw=.8, label="altitude (m, rel.)")
            ax[1].plot(T(lp), -lp["vz"], lw=.6, label="climb rate (m/s)")
        if mc is not None:
            for lab, keys, colr in (("pitch stick", ("pitch", "x"), "C0"), ("roll stick", ("roll", "y"), "C1"),
                                    ("throttle", ("throttle", "z"), "C2"), ("yaw stick", ("yaw", "r"), "C3")):
                k = next((k for k in keys if k in mc), None)
                if k is not None:
                    ax[2].plot(T(mc), mc[k], lw=.6, color=colr, label=lab)
        ax[2].set_ylabel("sticks")
        if act is not None:
            keys = [k for k in act if k.startswith("control[") or k.startswith("output[")][:8]
            for k in keys:
                ax[3].plot(T(act), act[k], lw=.5)
        ax[3].set_ylabel("actuators")
        if bat is not None:
            ax[4].plot(T(bat), bat["voltage_v"] if "voltage_v" in bat else bat.get("voltage_filtered_v"), lw=.8)
        ax[4].set_ylabel("battery V")
        if acc is not None:
            g = np.sqrt(acc["accelerometer_m_s2[0]"] ** 2 + acc["accelerometer_m_s2[1]"] ** 2 + acc["accelerometer_m_s2[2]"] ** 2) / 9.80665
            ax[5].plot(T(acc), g, lw=.4)
        ax[5].set_ylabel("accel |g|")
        ax[5].set_xlabel("seconds since arming")
        nav = st["nav_state"]
        ch = np.r_[0, np.where(np.diff(nav) != 0)[0] + 1]
        for i in ch:
            t = (st["timestamp"][i] - t_arm) / 1e6
            for x in ax:
                x.axvline(t, color="grey", lw=.4, ls=":")
            ax[0].text(t, 1.02, NAV.get(int(nav[i]), str(int(nav[i]))), transform=ax[0].get_xaxis_transform(),
                       fontsize=6, rotation=45)
        for x in ax:
            x.set_xlim(a, b)
            x.axvline(dur, color="red", lw=.6)
            x.grid(alpha=.3)
        ax[0].legend(fontsize=6, loc="lower left")
        ax[1].legend(fontsize=6, loc="lower left")
        ax[2].legend(fontsize=6, loc="lower left")
    axes[0, 0].set_title("Whole armed period (red line = disarm / end of log; grey dotted = mode change)", fontsize=9, pad=34)
    axes[0, 1].set_title("Last 60 s, enlarged", fontsize=9, pad=34)
    msgs = [(m.timestamp, m.log_level_str(), m.message) for m in u.logged_messages if m.log_level <= 4]
    lines = [f"{(t - t_arm) / 1e6:7.1f} s  {lvl:8s} {txt[:110]}" for t, lvl, txt in msgs[-15:]]
    fig.suptitle(title, fontsize=12, x=0.01, y=0.995, ha="left", fontweight="bold")
    fig.text(0.01, 0.005, "Last warnings/errors logged:\n" + ("\n".join(lines) if lines else "(none)"), fontsize=7,
             family="monospace", va="bottom")
    fig.subplots_adjust(bottom=0.16, top=0.90, hspace=0.12)
    fig.savefig(out_png, dpi=110)
    plt.close(fig)
    return dur


def workbook(path, cases):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    F, B = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
    wb = Workbook()
    ins = wb.active
    ins.title = "Instructions"
    text = [
        ("Independent crash relabelling", True), ("", False),
        ("You will see 72 real PX4 drone crash logs as plots (folder 'cases'). For each, decide from the log alone "
         "what most likely caused the crash. Work on your own and do not search for the logs online.", False), ("", False),
        ("Definitions", True),
        ("pilot = the crash would not have happened had the pilot acted correctly (inputs, mode choice, flying into "
         "something, ignoring a warning that left enough time).", False),
        ("aircraft = something in the aircraft, its software or configuration failed or behaved wrongly, and a "
         "reasonable pilot could not have prevented it.", False),
        ("environment = wind, obstacles or interference outside the pilot's and aircraft's control dominated.", False),
        ("can't tell = the plots do not show enough to decide, or causes contributed about equally.", False), ("", False),
        ("Each plot: left = whole flight, right = last 60 s enlarged. Panels: attitude with setpoints (dashed), altitude "
         "and climb rate, pilot sticks, actuator outputs, battery voltage, acceleration magnitude, flight modes (top "
         "labels), and the last logged warnings at the bottom. The red line is the end of the armed period; the crash "
         "may be before it, or may not be in the log at all ('no crash visible').", False), ("", False),
        ("Fill the yellow cells: verdict, main cause, confidence, and one line on what decided it. About 5 minutes per "
         "case. Do not open key_DO_NOT_OPEN.csv. Return only this workbook.", False)]
    for i, (t, bold) in enumerate(text, 1):
        c = ins.cell(row=i, column=1, value=t)
        c.font, c.alignment = (B if bold else F), Alignment(wrap_text=True, vertical="top")
    ins.column_dimensions["A"].width = 120
    ws = wb.create_sheet("Review")
    heads = ["Case", "Plot file", "Vehicle type", "Armed duration (s)", "Verdict", "Main cause", "Confidence", "What decided it"]
    for j, (h, w) in enumerate(zip(heads, [6, 16, 16, 12, 14, 30, 12, 60]), 1):
        c = ws.cell(row=1, column=j, value=h)
        c.font, c.fill = B, PatternFill("solid", fgColor="D9E1F2")
        ws.column_dimensions[c.column_letter].width = w
    for i, (n, vt, dur) in enumerate(cases, 2):
        for j, v in enumerate([n, f"case_{n:02d}.png", vt, round(dur, 1)], 1):
            ws.cell(row=i, column=j, value=v).font = F
        for j in range(5, 9):
            ws.cell(row=i, column=j).fill = PatternFill("solid", fgColor="FFFF00")
    last = len(cases) + 1
    for col, opts in (("E", VERDICTS), ("F", CAUSES), ("G", ["high", "medium", "low"])):
        dv = DataValidation(type="list", formula1='"' + ",".join(opts) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{last}")
    ws.freeze_panes = "B2"
    wb.save(path)


def main(root, out_dir):
    raw = open(os.path.join(root, "dbinfo.json"), "rb").read()
    db = {x["log_id"]: x for x in json.loads(gzip.decompress(raw) if raw[:2] == b"\x1f\x8b" else raw)}
    paths = pca.ulg_paths(root)
    os.makedirs(os.path.join(out_dir, "cases"), exist_ok=True)
    cases, key = [], []
    for n, (lid, lab) in enumerate(sample(), 1):
        vt = db.get(lid, {}).get("mav_type") or "unknown"
        dur = panel(paths[lid], os.path.join(out_dir, "cases", f"case_{n:02d}.png"), f"Case {n}   vehicle type: {vt}")
        cases.append((n, vt, dur))
        key.append([n, lid, lab])
        print(n, flush=True)
    workbook(os.path.join(out_dir, "relabel_workbook.xlsx"), cases)
    with open(os.path.join(out_dir, "key_DO_NOT_OPEN.csv"), "w", newline="") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(["case", "log_id", "uploader_label"])
        w.writerows(key)


if __name__ == "__main__":
    main(*sys.argv[1:3])
