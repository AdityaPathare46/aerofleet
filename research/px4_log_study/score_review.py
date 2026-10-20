"""Score the blind label review: do two reviewers agree with each other, and with the uploaders?

    python research/px4_log_study/score_review.py REVIEWER_A.xlsx REVIEWER_B.xlsx KEY_CSV [OUT_MD]
"""
import csv
import sys
from collections import Counter

from openpyxl import load_workbook

LABELS = ("pilot", "drone", "can't tell")


def read(path):
    ws = load_workbook(path, data_only=True)["Review"]
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[1]:
            out[row[1]] = {"verdict": (row[9] or "").strip().lower() or None, "cause": row[10], "confidence": row[11]}
    return out


def kappa(pairs):
    pairs = [(a, b) for a, b in pairs if a and b]
    n = len(pairs)
    if not n:
        return float("nan"), 0
    po = sum(a == b for a, b in pairs) / n
    ca, cb = Counter(a for a, _ in pairs), Counter(b for _, b in pairs)
    pe = sum(ca[k] * cb[k] for k in set(ca) | set(cb)) / n ** 2
    return ((po - pe) / (1 - pe) if pe < 1 else 1.0), n


def main(a_path, b_path, key_path, out=None):
    A, B = read(a_path), read(b_path)
    key = {r["log_id"]: r for r in csv.DictReader(open(key_path))}
    ids = list(key)
    L = ["# Crash label review: agreement", "",
         f"Rows with a verdict: reviewer A {sum(1 for i in ids if A.get(i, {}).get('verdict'))}, "
         f"reviewer B {sum(1 for i in ids if B.get(i, {}).get('verdict'))} of {len(ids)}.", ""]
    k, n = kappa([(A.get(i, {}).get("verdict"), B.get(i, {}).get("verdict")) for i in ids])
    L.append(f"- Reviewer A vs B, three classes (pilot / drone / can't tell): Cohen's kappa {k:.2f} on {n} crashes")
    for who, R in (("A", A), ("B", B)):
        v = [(R.get(i, {}).get("verdict"), key[i]["uploader_label"]) for i in ids]
        dec = [(x, y) for x, y in v if x in ("pilot", "drone")]
        ct = sum(1 for x, _ in v if x == "can't tell")
        agree = sum(x == y for x, y in dec)
        L.append(f"- Reviewer {who} vs uploader: agrees on {agree} of {len(dec)} decided crashes; "
                 f"'can't tell' on {ct}; kappa {kappa(dec)[0]:.2f}")
    both = [i for i in ids if A.get(i, {}).get("verdict") in ("pilot", "drone")
            and A.get(i, {}).get("verdict") == B.get(i, {}).get("verdict")]
    agree = sum(A[i]["verdict"] == key[i]["uploader_label"] for i in both)
    L.append(f"- Where both reviewers agree on pilot/drone ({len(both)} crashes), the uploader agrees on {agree}")
    L += ["", "| Uploader label | A: pilot | A: drone | A: can't tell |", "|---|---|---|---|"]
    for lab in ("pilot", "drone"):
        c = Counter(A.get(i, {}).get("verdict") for i in ids if key[i]["uploader_label"] == lab)
        L.append(f"| {lab} | {c['pilot']} | {c['drone']} | {c[LABELS[2]]} |")
    L += ["", "Read kappa as: below 0.4 weak, 0.4-0.6 moderate, 0.6-0.8 substantial, above 0.8 near-perfect."]
    text = "\n".join(L) + "\n"
    if out:
        open(out, "w").write(text)
    print(text)


if __name__ == "__main__":
    main(*sys.argv[1:5])
