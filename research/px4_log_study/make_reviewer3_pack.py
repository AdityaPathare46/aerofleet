"""Build the send pack for the third, non-author reviewer: two independent tasks in one folder and one zip.

  Task 1 - crash-moment check (third rater): the same blind 50-panel sheet the second reviewer labelled
           (crash_moment_sheet_blind.png; panels carry numbers only) and a blank form. ~30 minutes.
  Task 2 - crash relabelling: the blind pack from make_relabel_pack.py (72 plots + workbook). ~6 hours.
Never included: crash_moment_blind_key.json, either earlier reviewer's labels, key_DO_NOT_OPEN.csv. The panel key and
the earlier labels are public in the project repository, so the instructions ask the reviewer not to open it.

Output (git-ignored): research/px4_log_study/reviewer3_pack/ and reviewer3_pack.zip next to it.

    python research/px4_log_study/make_reviewer3_pack.py
"""
import os
import shutil
import zipfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
RESULTS = os.path.join(HERE, "..", "results", "px4_log_study")
OUT = os.path.join(HERE, "reviewer3_pack")
LABELS = ["correct", "too early", "too late", "no crash visible", "can't tell"]

README = """Thank you for reviewing. There are two separate tasks; please do them on your own, without discussing the cases
with anyone, and do not search for the logs online or open the project's code repository (it contains other
reviewers' answers).

TASK 1 - Crash-moment check (about 30 minutes)
  Files: Task1_crash_moment/crash_moment_sheet.png and Task1_crash_moment/crash_moment_form.xlsx
  The sheet has 50 panels from real drone crash logs. In each panel:
    black vertical line = the crash moment a computer picked
    blue line           = tilt angle in degrees (left axis; the dotted line is 60 degrees)
    orange line         = altitude in metres (right axis)
    x axis              = seconds relative to the black line
  For each panel choose one label in the form:
    correct           the black line is at the crash (loss of control / impact)
    too early         the crash happens clearly after the black line
    too late          the crash happens clearly before the black line
    no crash visible  nothing in the panel looks like a crash
    can't tell        the panel does not show enough to decide (e.g. empty panel)

TASK 2 - Crash relabelling (about 5 minutes per case, 72 cases)
  Files: Task2_relabel/cases/case_01.png ... case_72.png and Task2_relabel/relabel_workbook.xlsx
  Open the workbook's "Instructions" sheet first. For each case decide from the plots alone whether the pilot, the
  aircraft or the environment most likely caused the crash, the main cause, your confidence, and one line on what
  decided it.

Return only the two .xlsx files. Fill in the yellow cells; leave everything else as it is.
"""


def form(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Review"
    B, F = Font(name="Arial", size=10, bold=True), Font(name="Arial", size=10)
    for j, (h, w) in enumerate(zip(["Panel", "Label", "Note (optional)"], [8, 20, 60]), 1):
        c = ws.cell(row=1, column=j, value=h)
        c.font, c.fill = B, PatternFill("solid", fgColor="D9E1F2")
        ws.column_dimensions[c.column_letter].width = w
    for i in range(1, 51):
        ws.cell(row=i + 1, column=1, value=i).font = F
        for j in (2, 3):
            ws.cell(row=i + 1, column=j).fill = PatternFill("solid", fgColor="FFFF00")
    dv = DataValidation(type="list", formula1='"' + ",".join(LABELS) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("B2:B51")
    ws.freeze_panes = "A2"
    wb.save(path)


def main():
    src_pack = os.path.join(HERE, "relabel_pack")
    assert os.path.exists(os.path.join(src_pack, "relabel_workbook.xlsx")), "run make_relabel_pack.py first"
    shutil.rmtree(OUT, ignore_errors=True)
    t1, t2 = os.path.join(OUT, "Task1_crash_moment"), os.path.join(OUT, "Task2_relabel")
    os.makedirs(t1)
    shutil.copy(os.path.join(RESULTS, "crash_moment_sheet_blind.png"), os.path.join(t1, "crash_moment_sheet.png"))
    form(os.path.join(t1, "crash_moment_form.xlsx"))
    shutil.copytree(os.path.join(src_pack, "cases"), os.path.join(t2, "cases"),
                    ignore=shutil.ignore_patterns("._*"))
    shutil.copy(os.path.join(src_pack, "relabel_workbook.xlsx"), t2)
    open(os.path.join(OUT, "README_for_reviewer.txt"), "w").write(README)
    z = OUT + ".zip"
    with zipfile.ZipFile(z, "w", zipfile.ZIP_DEFLATED) as zf:
        for d, _, fs in os.walk(OUT):
            for f in fs:
                if not f.startswith("._"):
                    p = os.path.join(d, f)
                    zf.write(p, os.path.relpath(p, os.path.dirname(OUT)))
    names = zipfile.ZipFile(z).namelist()
    assert not any("key" in n.lower() or "blind_key" in n for n in names), "a key file got into the pack"
    print(f"{len(names)} files; {os.path.getsize(z) / 1e6:.1f} MB -> {z}")


if __name__ == "__main__":
    main()
