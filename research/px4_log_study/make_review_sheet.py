"""Build the blind label-review sheets for the PX4 crash benchmark.

60 crashes (30 the uploader blamed on the drone, 30 on the pilot), drawn from the crashes already analysed, in
random order. Each reviewer gets an identical workbook WITHOUT the uploader's label; the label goes to a separate
key file that reviewers must not open (it is git-ignored; rebuild it with this script). Seed 4.

    python research/px4_log_study/make_review_sheet.py DBINFO_JSON OUT_DIR
"""
import csv
import gzip
import html
import json
import os
import random
import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
RESULTS = os.path.join(HERE, "..", "results", "px4_log_study")
NAV = {0: "MANUAL", 1: "ALTCTL", 2: "POSCTL", 3: "MISSION", 4: "LOITER", 5: "RTL", 10: "ACRO", 12: "DESCEND",
       13: "TERMINATION", 14: "OFFBOARD", 15: "STABILIZED", 17: "TAKEOFF", 18: "LAND", 19: "FOLLOW", 20: "PRECLAND",
       21: "ORBIT", 22: "VTOL_TAKEOFF"}
SIGN_TEXT = {"D_BATT": "battery warning / low", "D_MOTOR": "motors at full output (last 10 s)",
             "D_CLIP": "accelerometer clipping", "D_GPS": "poor GPS in a GPS mode", "D_FD": "failure detector",
             "D_ERR": "autopilot error message", "D_CPU": "CPU overload", "D_FAILSAFE": "failsafe triggered"}
VERDICTS = ["pilot", "drone", "can't tell"]
CAUSES = ["pilot input / handling", "wrong mode or mode switch", "battery / power", "motor / ESC / propeller",
          "vibration / mounting", "GPS / compass / estimator", "radio or data link", "tuning / configuration",
          "software / firmware", "airframe / mechanical", "wind / environment", "not enough data"]
FONT, BOLD = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
INPUT = PatternFill("solid", fgColor="FFFF00")
HEAD = PatternFill("solid", fgColor="D9E1F2")
THIN = Border(bottom=Side(style="thin", color="BFBFBF"))


def load_rows():
    rows = []
    for name in ("pilot_vs_drone.jsonl", "holdout.jsonl"):
        for line in open(os.path.join(RESULTS, name)):
            r = json.loads(line)
            if r.get("group", "crash") == "crash" and r["label"] in ("drone", "pilot"):
                rows.append(r)
    return rows


def build(path, sample, notes):
    wb = Workbook()
    ins = wb.active
    ins.title = "Instructions"
    text = [
        ("PX4 crash label review", BOLD),
        ("", FONT),
        ("Why: the benchmark uses the uploaders' own 'pilot' or 'drone' labels. This review measures how reliable "
         "those labels are. Work on your own; do not discuss cases with the other reviewer until both are done.", FONT),
        ("", FONT),
        ("What to do for each row on the 'Review' sheet:", BOLD),
        ("1. Click 'Open plots' to see the log on PX4 Flight Review (attitude, motors, battery, GPS, RC inputs, modes, messages).", FONT),
        ("2. Read the facts in the row and the uploader's notes (written by the uploader; may be empty or wrong).", FONT),
        ("3. Fill the yellow cells: Your verdict, Main cause, Confidence, and a short note on what decided it.", FONT),
        ("", FONT),
        ("Definitions:", BOLD),
        ("pilot = the crash would not have happened had the pilot acted correctly (inputs, mode choice, flying into something, "
         "ignoring a warning that gave enough time).", FONT),
        ("drone = something in the aircraft or its software/configuration failed or behaved wrongly, and a reasonable pilot could not prevent it.", FONT),
        ("can't tell = the log does not show enough to decide, or both contributed about equally.", FONT),
        ("", FONT),
        ("Rules: about 5 minutes per crash; spend no more than 10. Do not look up the log's rating on Flight Review "
         "before deciding. Do not open key_DO_NOT_OPEN.csv.", FONT),
        ("", FONT),
        ("Example of a filled row (not a real case):", BOLD),
    ]
    for i, (t, f) in enumerate(text, 1):
        c = ins.cell(row=i, column=1, value=t)
        c.font = f
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ex_row = len(text) + 1
    for j, (h, v) in enumerate([("Your verdict", "drone"), ("Main cause", "motor / ESC / propeller"),
                                ("Confidence", "medium"),
                                ("Your note", "Motor 3 output at 100% from 41.2 s while the others drop; vehicle yaws then flips. Pilot inputs normal.")]):
        ins.cell(row=ex_row, column=1 + 2 * j, value=h).font = BOLD
        ins.cell(row=ex_row, column=2 + 2 * j, value=v).font = FONT
    ins.cell(row=ex_row + 2, column=1, value="Progress (rows with a verdict):").font = BOLD
    p = ins.cell(row=ex_row + 2, column=2, value=f"=COUNTA(Review!J2:J{len(sample) + 1})&\" of {len(sample)}\"")
    p.font = FONT
    ins.column_dimensions["A"].width = 110
    for col in "BCDEFGH":
        ins.column_dimensions[col].width = 22

    ws = wb.create_sheet("Review")
    heads = ["#", "Log ID", "Open plots", "Vehicle", "Flight time to crash (s)", "Mode at crash", "Rotary wing",
             "Signs that fired (rule checks)", "Uploader notes", "Your verdict", "Main cause", "Confidence", "Your note"]
    widths = [5, 38, 12, 16, 12, 14, 8, 34, 60, 14, 26, 12, 50]
    for j, (h, w) in enumerate(zip(heads, widths), 1):
        c = ws.cell(row=1, column=j, value=h)
        c.font, c.fill = BOLD, HEAD
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[c.column_letter].width = w
    ws.cell(row=1, column=8).comment = Comment(
        "Fixed checks from pilot_vs_drone.py. They fire on many normal flights too (see holdout.md); treat them as "
        "pointers, not answers.", "AeroFleet")
    for i, r in enumerate(sample, 2):
        signs = [SIGN_TEXT[k] for k, v in r["evidence"].items() if v]
        vals = [i - 1, r["log_id"], None, r["mav_type"], r["flight_s"], NAV.get(r["mode_at_crash"], str(r["mode_at_crash"])),
                "yes" if r["rotary"] else "no", "; ".join(signs) or "none", notes.get(r["log_id"], "")]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = FONT
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = THIN
        link = ws.cell(row=i, column=3, value="Open plots")
        link.hyperlink = f"https://review.px4.io/plot_app?log={r['log_id']}"
        link.font = Font(name="Arial", size=10, color="0563C1", underline="single")
        for j in range(10, 14):
            c = ws.cell(row=i, column=j)
            c.fill, c.font, c.border = INPUT, FONT, THIN
            c.alignment = Alignment(wrap_text=True, vertical="top")
    n = len(sample) + 1
    for col, opts in (("J", VERDICTS), ("K", CAUSES), ("L", ["high", "medium", "low"])):
        dv = DataValidation(type="list", formula1='"' + ",".join(opts) + '"', allow_blank=True)
        dv.error, dv.errorTitle = "Pick a value from the list.", "Not in the list"
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{n}")
    ws.freeze_panes = "C2"
    wb.save(path)


def main(dbinfo, out_dir, n_each=30, seed=4):
    os.makedirs(out_dir, exist_ok=True)
    raw = open(dbinfo, "rb").read()
    db = json.loads(gzip.decompress(raw) if raw[:2] == b"\x1f\x8b" else raw)
    notes = {}
    for x in db:
        t = " | ".join(s.strip() for s in ((x.get("description") or ""), (x.get("feedback") or "")) if s.strip())
        notes[x["log_id"]] = " ".join(html.unescape(t).split())[:600]
    rows = load_rows()
    rng = random.Random(seed)
    sample = (rng.sample([r for r in rows if r["label"] == "drone"], n_each)
              + rng.sample([r for r in rows if r["label"] == "pilot"], n_each))
    rng.shuffle(sample)
    for who in ("A", "B"):
        build(os.path.join(out_dir, f"crash_review_reviewer_{who}.xlsx"), sample, notes)
    with open(os.path.join(out_dir, "key_DO_NOT_OPEN.csv"), "w", newline="") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(["#", "log_id", "uploader_label", "rule_v1_verdict"])
        for i, r in enumerate(sample, 1):
            w.writerow([i, r["log_id"], r["label"], r["verdict"]])
    print(f"{len(sample)} crashes; files in {out_dir}")


if __name__ == "__main__":
    main(*sys.argv[1:3])
