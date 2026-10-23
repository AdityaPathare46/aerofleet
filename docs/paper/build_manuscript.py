"""Build the manuscript (.docx) from text kept here and figures from research/px4_log_study/paper_figures.py.

    python3 docs/paper/build_manuscript.py
Every number in the text comes from a committed file under research/results/; see docs/paper/README.md.
"""
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.join(HERE, "figures")

TITLE = ("Flight logs reveal which drones are at risk, not why a flight crashed: "
         "pre-declared tests on 479,547 public PX4 flight logs")

HIGHLIGHTS = [
    "Uploader-labelled crashes from 479,547 public PX4 logs form a real-flight benchmark",
    "No rule or model separated pilot-caused from aircraft-caused crashes above chance",
    "The flight that crashed was no rougher than the same drone's earlier flights",
    "Drones that later crashed vibrated more than matched drones on ordinary flights",
    "Small language-model review panels lost to a fixed rule and were easy to steer",
]

ABSTRACT = (
    "Operators, insurers and investigators increasingly expect flight logs to explain why a small drone crashed. "
    "We tested that expectation on the public PX4 Flight Review archive (479,547 logs; 590 crashes labelled by their "
    "uploaders as pilot- or aircraft-caused). Every test was declared, with its threshold, in time-stamped version "
    "control before it was run, and decisive tests were repeated on held-out data. Safety settings did not separate "
    "crash logs from well-rated flights (a protection disabled in 51% vs 46%, p = 0.63). A fixed attribution rule "
    "reached a balanced accuracy of 0.47 on 112 held-out crashes, and a cross-validated model stayed between 0.42 and "
    "0.52 at every cut-off before the crash. Checks that appeared to find the cause named by forum experts mostly "
    "fired on the crash itself (13 of 15 tuning faults found with the crash included, 2 of 10 without its last 10 s). "
    "Crash flights were distinguishable from other drones' flights (AUC about 0.86 even 30 s before the crash) but not "
    "from the same drone's earlier flights (vibration rank 0.48, p = 0.67). Against matched drones without a recorded "
    "crash, drones that later crashed vibrated more on ordinary flights (mean rank 0.67, 95% CI 0.57-0.77, p = 0.002, "
    "60 matched sets). Small language-model review panels lost to a fixed rule and could be steered by one compromised "
    "reviewer in 45-56 of 60 cases. Flight logs therefore support ranking airframes for inspection, not attributing an "
    "individual crash. We release the selection, code and labels.")

KEYWORDS = ["Unmanned aircraft systems", "Flight log analysis", "Crash investigation", "Fault diagnosis",
            "Data leakage", "Pre-registration"]

SECTIONS = [
    ("1. Introduction", [
        "Small unmanned aircraft are now flown commercially for survey, inspection, agriculture and delivery, and the "
        "rules that govern them place weight on records. In the European Union and the United Kingdom an operator in "
        "the specific category must keep records of each operation and of maintenance for at least three years "
        "[12, 19]; India requires an accident to be reported within 48 hours [13], and the United States within ten "
        "days [21]. When something goes wrong, the onboard "
        "flight log is the richest record available, and a growing set of tools, from open-source log viewers to "
        "commercial fleet platforms and recent machine-learning prototypes, promise to turn it into an answer: what "
        "failed, and whether the pilot or the aircraft was at fault.",
        "Whether logs can actually deliver that answer has not, to our knowledge, been tested on real crashes at "
        "scale. Public fault datasets are dominated by simulation or by faults injected on purpose [1-3], and accident "
        "statistics classify outcomes such as loss of control rather than causes visible in a log [8, 11]. Two pitfalls "
        "make optimistic results easy to obtain. First, a log that contains the crash also contains its consequences: "
        "saturated motors, extreme attitudes and estimator errors, which a detector can learn instead of the cause, a "
        "form of data leakage [16, 17]. Second, comparing crash flights with other operators' good flights confounds "
        "the flight with the aircraft and its user.",
        "This paper asks four questions of the largest public archive of real autopilot logs, PX4 Flight Review: "
        "(Q1) do configuration settings differ between crash and normal flights; (Q2) can a log separate "
        "pilot-caused from aircraft-caused crashes; (Q3) does the flight that crashed show warning signs relative to "
        "the same drone's history, or are differences between drones; and (Q4) do expert diagnoses agree with uploader "
        "labels and with automatic checks. Because language-model assistants are now proposed for exactly this task, "
        "we also report a controlled test of small local models as review panels.",
        "Our contributions are: (i) a real-flight crash benchmark built from public logs, with a check of its labels "
        "against expert diagnoses; (ii) a set of tests, each declared with its threshold before it was run and "
        "repeated on held-out data where it decided anything; (iii) evidence that crash attribution from logs is not "
        "achievable with these labels and features, and that apparent diagnostic accuracy can come from the crash "
        "itself; (iv) a positive, practically usable finding: drones that later crash are rougher on their ordinary "
        "flights than matched drones, which supports fleet-level inspection ranking but not per-flight prediction.",
    ]),
    ("2. Related work", [
        "Datasets. ALFA provides 47 fixed-wing flights with injected engine and control-surface failures [1]. RflyMAD "
        "covers 5,629 multicopter fault cases, of which 497 are real flights and the rest software- or hardware-in-the-"
        "loop simulations [2]. UAV-SEAD offers about 1,396 real PX4 logs annotated for state-estimation anomalies [3], "
        "used by later anomaly-mining work [4]; a recent testbed quantifies the gap between simulated and real PX4 logs "
        "[5]. These resources target anomaly or fault detection; none labels real crashes by responsibility, and none "
        "separates the pre-crash window from the crash.",
        "Diagnosis and causality. Causal models built from autopilot documentation have been used to find actual causes "
        "of failures in ArduPilot logs [6]. A study of 569 bugs in PX4 and ArduPilot found parameter handling among the "
        "root-cause types of drone-specific bugs [7]. Language-model agents have been applied to simulation testing and "
        "log analytics [9], and to classifying problems in drone log messages for forensics [10]. Community tools for "
        "automated log diagnosis exist but are evaluated mainly on simulated faults.",
        "Accident analysis. Analyses of civil drone occurrences find equipment problems and loss of control prominent "
        "[8]; human-factors classifications of remotely piloted aircraft mishaps attribute a large share of causal "
        "factors to people but note substantial rater disagreement [18, 22]. Our forum analysis adds expert "
        "diagnoses of individual real crashes, linked to their logs.",
        "Methodology. Leakage, where information that would not be available at prediction time enters training or "
        "evaluation, is a known source of irreproducible machine-learning results [16, 17]; label noise likewise "
        "biases evaluation [20]. Pre-registration separates confirmatory from exploratory analysis [15]. We apply both "
        "ideas to drone logs.",
    ]),
    ("3. Data", [
        "Archive. We downloaded the public index of PX4 Flight Review (479,547 logs on 20 October 2026; 442,929 from "
        "real hardware rather than simulation, from 38,320 distinct vehicle identifiers). Uploaders may rate a log; 590 "
        "carry a crash rating, 402 as a software or hardware cause and 188 as a pilot cause. Logs were read with "
        "pyulog [14]. Only public logs were used; the analysis stores derived features and log identifiers, not the "
        "logs themselves.",
        "Samples. Table 1 lists the samples. Crash samples used real hardware, PX4 1.10 or newer and at least 5 s of "
        "flight unless stated. Normal flights were logs rated good or great, matched on vehicle type and drawn from "
        "vehicles without a crash rating. For the within-drone test every crash with a vehicle identifier and at least "
        "three earlier non-crash flights of 30 s or more was eligible (101 crashes). For the between-drone test each "
        "crashed drone was matched with up to two drones without any crash-rated log that shared the flight-controller "
        "hardware, PX4 major.minor version and vehicle type (66 matched sets, 124 control drones).",
        "Expert labels. We searched the PX4 discussion forum for threads that link a Flight Review log and mention a "
        "crash (443 threads, 410 with a log link). For the 161 crash threads with a reply from PX4 staff, the "
        "development team or an experienced member, we recorded the expert's conclusion, the cause named, who gave it "
        "and whether the original poster confirmed it, with the sentence stating the diagnosis. Recording was "
        "AI-assisted and reviewed by an author; 123 of 124 quoted sentences were found verbatim in the thread by script. "
        "An independent second recording is in progress, and the error rate of the recording is not yet measured.",
    ]),
    ("4. Methods", [
        "Declaration. Each test, its features, thresholds and pass mark were committed to a public version-control "
        "repository before the data it used were downloaded or analysed; the commit identifiers are listed with the "
        "results. One amendment was made before any analysis: the firmware filter of the within-drone test was removed "
        "because it left only 12 cases. One coding error was corrected after a first run (logs with a CPU-load record "
        "lacking the load field were being skipped); no threshold changed.",
        "Crash moment. For rotary-wing aircraft the crash moment t_c is the first time after take-off at which the tilt "
        "exceeds 60 degrees; otherwise, and for fixed-wing aircraft, it is the end of the armed period. Normal flights "
        "use the end of the armed period. Analyses that exclude the crash stop the window m seconds before t_c.",
        "Features. Per window: root-mean-square rate-tracking error (roll and pitch setpoint minus measured rate), "
        "95th percentile of tilt and of accelerometer vibration, accelerometer clipping per second, share of time with "
        "any motor at 95% or more, lowest battery percentage, relative voltage sag, share of time with poor GPS "
        "(horizontal accuracy over 5 m or fewer than 6 satellites), any failsafe, share of time in a pilot-flown mode, "
        "and the 95th percentile of stick deflection. Older firmware lacks some messages; there vibration falls back to "
        "the estimator's vibration estimate and measured rate to the attitude message, and comparisons are made only "
        "between flights measured the same way.",
        "Tests. Q1 compared disabled protections (relative to the firmware default record) between crash and normal "
        "logs with Fisher's exact test. Q2 used a fixed rule (any of eight aircraft-fault signs implies 'aircraft', "
        "otherwise a pilot-flown mode at t_c implies 'pilot') with a declared pass mark of balanced accuracy 0.70, and "
        "an L2-regularised logistic regression evaluated by stratified five-fold cross-validation repeated five times. "
        "The leakage test (Q3a) declared that crash-versus-normal AUC would fall by at least 0.10 between m = 0 and "
        "m = 10 s. Q3b ranked each crash flight among the same drone's earlier flights; Q3c ranked each crashed drone, "
        "by the median over three ordinary flights, among its matched controls. Both used the scaled rank (0 lowest, "
        "1 highest), whose null mean is 0.5, with one-sided p from 100,000 draws of the exact null and alpha 0.025 per "
        "hypothesis (two primary features). Q4 compared expert conclusions with uploader labels and tested whether "
        "fixed checks fired on logs whose expert named the matching fault, at most 10% of normal flights.",
        "Language-model panels. Sixty generated cases with known causes were presented to five reviewers, each seeing "
        "only numbered facts from one evidence area, with citations checked by code and a fixed aggregation rule; a "
        "single agent shown all facts and a fixed no-AI rule served as comparators. Four local models of 3-4 billion "
        "parameters were tested at temperature 0. Robustness was measured by replacing one reviewer with an adversary "
        "arguing for a chosen wrong cause. Hostile text was also planted in five uploader-controlled fields to check "
        "whether it could reach a prompt.",
    ]),
    ("5. Results", []),
]

RESULTS = [
    ("5.1 Configuration settings (Q1)",
     "Of 57 crash logs and 113 normal logs that carried a firmware-default record, 29 (51%) and 52 (46%) had at least "
     "one protection disabled relative to the default (Fisher p = 0.63). Pilot-blamed crashes had a protection disabled "
     "more often than aircraft-blamed ones (14 of 20 vs 15 of 37), a contrast too small to test. None of 100 crash logs "
     "changed a safety-critical or flight-limit setting during the flight; the only in-flight change observed was a "
     "flight counter."),
    ("5.2 Attributing a crash to the pilot or the aircraft (Q2)",
     "The fixed rule reached a balanced accuracy of 0.57 on 93 development crashes and 0.47 on 112 held-out crashes, "
     "against a declared pass mark of 0.70, and called 39% of 143 normal flights an aircraft fault (Table 2). The "
     "aircraft-fault signs fired at similar rates whichever party the uploader blamed: motor saturation in the last 10 s "
     "in 23% of aircraft-blamed and 29% of pilot-blamed crashes (1% of normal flights), battery warnings in 15% and 18% "
     "(3%). Error messages (19%, 10%; normal 20%) and failsafes (18%, 8%; normal 13%) were no more frequent in crashes "
     "than in normal flights. Two signals noticed after the first run did not replicate on held-out crashes "
     "(accelerometer clipping, p = 0.045; manual mode at the crash, p = 0.165; declared alpha 0.025). The "
     "cross-validated model stayed between 0.42 and 0.52 at every cut-off (Fig. 1)."),
    ("5.3 The crash window (Q3a)",
     "Crash flights were separated from normal flights with a cross-validated AUC of 0.88 when the whole flight was used "
     "and 0.86 when the last 30 s were removed (Fig. 1); the declared drop of 0.10 by 10 s was not observed (0.01). "
     "Leakage is nevertheless present at the level of individual checks: the rate-tracking feature fell from 0.81 to "
     "0.69 as the crash was excluded, and a fixed tuning check that found 13 of 15 expert-diagnosed tuning faults with "
     "the crash included found 2 of 10 once the last 10 s were removed (Section 5.5)."),
    ("5.4 Within the same drone, and between drones (Q3b, Q3c)",
     "Relative to the same drone's earlier flights, the crash flight was neither more vibrating (mean scaled rank 0.48, "
     "95% CI 0.38-0.58, p = 0.67, 61 drones) nor worse at tracking its commands (0.50, 0.38-0.63, p = 0.47, 46 drones; "
     "Fig. 2). Relative to matched drones without a recorded crash, drones that later crashed vibrated more on their "
     "ordinary pre-crash flights (0.67, 0.57-0.77, p = 0.0023, 60 sets), while tracking error did not differ "
     "significantly (0.57, 0.44-0.70, p = 0.15, 47 sets). The vibration difference was clearest for aircraft-blamed "
     "crashes (0.67, p = 0.005, 46 sets) and similar in size but not significant for pilot-blamed ones (0.64, p = 0.16, "
     "14 sets). The crash-versus-normal separation of Section 5.3 is therefore a property of the drone, not a build-up "
     "within the flight."),
    ("5.5 Expert diagnoses (Q4)",
     "Of 148 forum threads describing a real crash, the expert attributed it to the aircraft in 88 (59%), to the pilot "
     "in 9 (6%), to both in 2, and was unclear or gave no diagnosis in 49 (33%). Tuning or configuration was the cause "
     "named most often (29 of 88; 15 of the 47 confirmed by the poster), followed by GPS, compass or estimator problems "
     "(16) and motors, ESCs or propellers (14) (Fig. 3). Of 14 threads whose log was rated an aircraft crash by its "
     "uploader, the expert agreed in 9 and was unclear in 5; none attributed it to the pilot. Only 35 of the 70 logs "
     "behind aircraft diagnoses could still be downloaded. With the window ending 10 s before the crash, no fixed check "
     "met the declared criterion: tuning 2 of 10 (5% of normal flights), vibration 3 of 3 (11%), GPS or compass 0 of 3, "
     "motor 1 of 4, battery 0 of 2."),
    ("5.6 Language-model review panels",
     "The no-AI rule was right in 45 of 60 generated cases and wrong in none. No model's panel matched it (Table 3): "
     "right/wrong counts were 3/36, 20/19, 10/6 and 17/38 for gemma3:4b, phi4-mini, qwen2.5:3b and qwen3:4b. A single "
     "compromised reviewer changed the panel's answer in 49, 56, 54 and 45 of 60 cases respectively. None of 40 planted "
     "hostile strings reached a model prompt, because the pipeline passes only numbers extracted by code."),
]

DISCUSSION = [
    ("6. Discussion", [
        "Attribution. Three independent attempts, a declared rule, held-out replication and a cross-validated model, "
        "found no usable separation between pilot- and aircraft-caused crashes. Uploader labels are noisy, but the "
        "expert threads suggest they are not random: where an expert reached a verdict on an aircraft-rated crash, it "
        "agreed. The more likely explanation is that the consequences of a pilot error and of an aircraft fault look "
        "alike in the signals a log records, as forum experts' frequent 'unclear' verdicts also indicate. Tools that "
        "issue automatic responsibility verdicts from logs should be treated with caution.",
        "Leakage. The model-level leakage we expected was absent, yet single checks showed it clearly. Evaluations of "
        "log-based diagnosis should report results with the crash window removed, and should compare against the same "
        "aircraft's history before claiming that a signal anticipates a crash.",
        "Drone-level risk. The strongest positive result concerns airframes, not flights: drones that later crashed "
        "vibrated more than comparable drones on ordinary flights, and expert diagnoses most often named tuning, "
        "configuration and propulsion problems that raise vibration. This supports ranking a fleet's drones by "
        "vibration and inspecting the roughest, and it does not support a per-flight crash warning. We implemented "
        "such a ranking in an open fleet-assurance tool, flagging drones above their group's upper quartile and only "
        "against drones measured the same way, with the explicit caveat that a flag is a prompt to inspect.",
        "Language models. On this task small local models added errors and an attack surface rather than accuracy. "
        "The result concerns 3-4 billion-parameter models that run on ordinary laptops; larger models were not tested.",
    ]),
    ("7. Threats to validity", [
        "Labels are the uploaders' own and normal flights are self-rated; controls without a crash-rated log may still "
        "have crashed, which would bias the between-drone test towards no difference. Logs on the archive are "
        "self-selected and PX4-only; ArduPilot fleets were not studied. The crash moment is estimated. Matching used "
        "hardware, firmware and vehicle type, not frames, propellers or missions. Sample sizes for some comparisons are "
        "small (14 pilot-blamed matched sets; 3 vibration-diagnosed logs). Forum labels were recorded with AI assistance "
        "by one recorder, and half of the expert-diagnosed logs are no longer public. The language-model cases are "
        "generated, not real incidents.",
    ]),
    ("8. Conclusion", [
        "On real, public crash logs, settings did not distinguish crashes, simple rules and models could not say "
        "whether the pilot or the aircraft was responsible, and the flight that crashed looked like the same drone's "
        "earlier flights. What logs did reveal was which drones were rougher than their peers. We suggest that log "
        "analysis for operators be aimed at fleet-level inspection ranking and at presenting evidence for a human "
        "investigator, and that claims of automatic crash attribution be evaluated with the crash window removed and "
        "against each aircraft's own history.",
    ]),
]

TABLE1 = [("Sample", "Crashes", "Comparison", "Used for"),
          ("Settings sample", "100", "477 normal (header only)", "Q1"),
          ("Development crashes", "93 analysed", "-", "Q2 rule"),
          ("Held-out crashes", "112 analysed", "143 normal (full)", "Q2 replication, signs"),
          ("Leak-curve set", "148 at m = 0", "139 normal", "Q2 model, Q3a"),
          ("Within-drone set", "101 selected (61 / 46 used)", "475 earlier flights of the same drones", "Q3b"),
          ("Between-drone set", "66 matched sets (60 / 47 used)", "124 matched drones, 372 flights", "Q3c"),
          ("Forum threads", "148 real crashes", "35 logs downloadable", "Q4")]

TABLE2 = [("Test (declared pass mark)", "Result", "Outcome"),
          ("Protections disabled, crash vs normal", "51% vs 46%, p = 0.63", "No difference"),
          ("Attribution rule, development (BA >= 0.70)", "BA 0.57, 35% wrong", "Failed"),
          ("Attribution rule, held-out", "BA 0.47, 47% wrong; 39% of normal flagged", "Failed"),
          ("Clipping in aircraft-blamed crashes (p < 0.025)", "24% vs 8%, p = 0.045", "Not supported"),
          ("Manual mode in pilot-blamed crashes (p < 0.025)", "18% vs 10%, p = 0.165", "Not supported"),
          ("Leak: AUC drop 0 to 10 s (>= 0.10)", "0.88 to 0.87 (drop 0.01)", "Not supported"),
          ("Within-drone vibration (p < 0.025)", "rank 0.48, p = 0.67", "Not supported"),
          ("Within-drone tracking (p < 0.025)", "rank 0.50, p = 0.47", "Not supported"),
          ("Between-drone vibration (p < 0.025)", "rank 0.67, p = 0.0023", "Supported"),
          ("Between-drone tracking (p < 0.025)", "rank 0.57, p = 0.15", "Not supported"),
          ("Expert-fault checks, 10 s before (>= 50% / <= 10%)", "best: vibration 3/3 vs 11%", "None met")]

TABLE3 = [("Model", "Single agent right/wrong", "Panel right/wrong", "Steered by one reviewer"),
          ("No-AI rule", "45 / 0", "-", "-"),
          ("gemma3:4b", "10 / 0", "3 / 36", "49 of 60"),
          ("phi4-mini", "10 / 0", "20 / 19", "56 of 60"),
          ("qwen2.5:3b", "9 / 44", "10 / 6", "54 of 60"),
          ("qwen3:4b", "0 / 41", "17 / 38", "45 of 60")]

REFERENCES = [
    "Keipour A, Mousaei M, Scherer S. ALFA: A dataset for UAV fault and anomaly detection. Int J Robot Res 2021;40(2-3):515-20.",
    "RflyMAD: A dataset for multicopter fault detection and health assessment. Int J Robot Res 2025; arXiv:2311.11340. [authors and pages to verify]",
    "Kabaoglu A, Sariel S. UAV-SEAD: State estimation anomaly dataset for UAVs. arXiv:2602.13900; 2026.",
    "AeroTSBoost: Temporal-statistical boosting for real-world UAV telemetry anomaly mining. arXiv:2605.25639; 2026. [authors to verify]",
    "A UAV testbed for diagnosing hardware vulnerabilities: Quantifying sim-to-real discrepancies in PX4 flight logs. Sensors 2026;26(10):3188. [authors to verify]",
    "Zibaei E, Borth R. Building causal models for finding actual causes of unmanned aerial vehicle failures. Front Robot AI 2024;11:1123762.",
    "Wang D, Li S, Xiao G, Liu Y, Sui Y. An exploratory study of autopilot software bugs in unmanned aerial vehicles. In: Proc ESEC/FSE 2021. p. 20-31. [pages to verify]",
    "Wild G, Murray J, Baxter G. Exploring civil drone accidents and incidents to help prevent potential air disasters. Aerospace 2016;3(3):22.",
    "LLM-agents driven automated simulation testing and analysis of small uncrewed aerial systems. arXiv:2501.11864; 2025. [authors to verify]",
    "Silalahi S, Ahmad T, Studiawan H, Breitinger F. DroPTC: problem-type classification of drone log messages for forensic analysis. DFRWS EU 2026. [title to verify]",
    "Air Accidents Investigation Branch. Annual Safety Review 2023. Farnborough: AAIB; 2024.",
    "Commission Implementing Regulation (EU) 2019/947 on the rules and procedures for the operation of unmanned aircraft. OJ L 152, 11.6.2019.",
    "Ministry of Civil Aviation, Government of India. The Drone Rules, 2021. G.S.R. 589(E), 25 August 2021.",
    "PX4 Development Team. pyulog: Python library for parsing ULog files; and PX4 Flight Review. https://github.com/PX4/pyulog; https://review.px4.io.",
    "Nosek BA, Ebersole CR, DeHaven AC, Mellor DT. The preregistration revolution. Proc Natl Acad Sci USA 2018;115(11):2600-6.",
    "Kapoor S, Narayanan A. Leakage and the reproducibility crisis in machine-learning-based science. Patterns 2023;4(9):100804.",
    "Kaufman S, Rosset S, Perlich C, Stitelman O. Leakage in data mining: Formulation, detection, and avoidance. ACM Trans Knowl Discov Data 2012;6(4):15.",
    "Wiegmann DA, Shappell SA. A human error approach to aviation accident analysis: The human factors analysis and classification system. Aldershot: Ashgate; 2003.",
    "European Union Aviation Safety Agency. Easy Access Rules for Unmanned Aircraft Systems. Cologne: EASA; latest revision.",
    "Frenay B, Verleysen M. Classification in the presence of label noise: A survey. IEEE Trans Neural Netw Learn Syst 2014;25(5):845-69.",
    "Federal Aviation Administration. 14 CFR Part 107, Small unmanned aircraft systems; section 107.9, accident reporting.",
    "Tvaryanas AP, Thompson WT, Constable SH. Human factors in remotely piloted aircraft operations: HFACS analysis of 221 mishaps over 10 years. Aviat Space Environ Med 2006;77(7):724-32.",
]

DECLARATIONS = [
    ("Data and code availability",
     "Selection scripts, analysis code, declared hypotheses with their commit history, derived features and the expert "
     "labels (without forum text) are available at [repository URL]. Raw logs are public on PX4 Flight Review and are "
     "referenced by identifier."),
    ("Declaration of generative AI and AI-assisted technologies in the writing process",
     "During the preparation of this work the authors used an AI assistant (Claude, Anthropic) to write analysis code, "
     "draft text and assist the recording of forum diagnoses. The authors reviewed and edited the content and take full "
     "responsibility for the publication."),
    ("Declaration of competing interest", "The authors declare no competing interests. [confirm]"),
    ("CRediT authorship contribution statement", "[Author 1]: Conceptualization, Methodology, Software, Writing - original draft. "
     "[Author 2]: Investigation, Validation. [complete]"),
    ("Funding", "This research did not receive any specific grant. [confirm]"),
]


def para(doc, text, bold=False, italic=False, size=11, align=None, space_after=6):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold, r.italic, r.font.size = bold, italic, Pt(size)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.5
    if align:
        p.alignment = align
    return p


def table(doc, rows, caption):
    para(doc, caption, bold=True, size=10, space_after=3)
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            c = t.cell(i, j)
            c.text = cell
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
                    r.bold = i == 0
    doc.add_paragraph()


def figure(doc, name, caption, width=14):
    doc.add_picture(os.path.join(FIG, name), width=Cm(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    para(doc, caption, size=9.5, italic=True)


def main():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(11)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.5)
    para(doc, TITLE, bold=True, size=15, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "[Author names and affiliations]  ·  Corresponding author: [name, e-mail]", size=10,
         align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "Highlights", bold=True, size=12)
    for h in HIGHLIGHTS:
        doc.add_paragraph(h, style="List Bullet")
    para(doc, "Abstract", bold=True, size=12)
    para(doc, ABSTRACT)
    para(doc, "Keywords: " + "; ".join(KEYWORDS), italic=True, size=10)
    for title, body in SECTIONS:
        para(doc, title, bold=True, size=12, space_after=4)
        for b in body:
            para(doc, b)
        if title.startswith("3."):
            table(doc, TABLE1, "Table 1. Samples drawn from the PX4 Flight Review archive.")
        if title.startswith("5."):
            for sub, text in RESULTS:
                para(doc, sub, bold=True, italic=True, size=11, space_after=3)
                para(doc, text)
                if sub.startswith("5.2"):
                    table(doc, TABLE2, "Table 2. Declared tests and their outcomes.")
                    figure(doc, "fig_leak_curve.png", "Fig. 1. Cross-validated AUC as the end of the window moves away from "
                           "the crash moment. Crash-versus-normal separation does not depend on the crash window; "
                           "pilot-versus-aircraft attribution stays at chance. Bars: SD over five repeats of five-fold CV.")
                if sub.startswith("5.4"):
                    figure(doc, "fig_within_between.png", "Fig. 2. The crashed drone's rank against its own earlier "
                           "flights (top) and against matched drones without a recorded crash (bottom).")
                if sub.startswith("5.5"):
                    figure(doc, "fig_forum_causes.png", "Fig. 3. Causes named by experts in PX4 forum threads that "
                           "attributed a crash to the aircraft.")
                if sub.startswith("5.6"):
                    table(doc, TABLE3, "Table 3. Language-model review panels on 60 generated cases (counts of cases).")
    for title, body in DISCUSSION:
        para(doc, title, bold=True, size=12, space_after=4)
        for b in body:
            para(doc, b)
    for title, text in DECLARATIONS:
        para(doc, title, bold=True, size=11, space_after=3)
        para(doc, text, size=10)
    para(doc, "References", bold=True, size=12)
    for i, ref in enumerate(REFERENCES, 1):
        para(doc, f"[{i}] {ref}", size=9.5, space_after=2)
    out = os.path.join(HERE, "AeroFleet_flight_log_crash_study_manuscript.docx")
    doc.save(out)
    words = sum(len(t.split()) for t in [ABSTRACT] + [b for _, bs in SECTIONS + DISCUSSION for b in bs] + [t for _, t in RESULTS])
    print(out, "main text words:", words)


if __name__ == "__main__":
    main()
