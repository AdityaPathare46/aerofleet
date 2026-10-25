"""Build a Word version of the manuscript from latex/main.tex (pandoc + a little clean-up).

    cd docs/paper/latex && tectonic -X compile --keep-intermediates main.tex   # main.aux gives the numbers
    python docs/paper/make_docx.py

The LaTeX stays the master copy. This makes submission/manuscript.docx with the same title page, numbered
sections, tables and figures (numbers taken from main.aux, so they match the PDF), and numbered references in the
Elsevier style (elsevier-with-titles.csl).
"""
import os
import re
import subprocess
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
TEX, AUX = os.path.join(HERE, "latex", "main.tex"), os.path.join(HERE, "latex", "main.aux")
OUT = os.path.join(HERE, "submission", "manuscript.docx")
TMP = os.path.join(HERE, "latex", "_docx_input.tex")


def labels():
    out = {}
    content = open(AUX).read()
    for m in re.finditer(r"\\newlabel\{([^}]+)\}\{", content):
        k = m.group(1)
        i = m.end()
        if i < len(content) and content[i] == "{":
            depth = 1
            j = i + 1
            while j < len(content) and depth > 0:
                if content[j] == "{":
                    depth += 1
                elif content[j] == "}":
                    depth -= 1
                j += 1
            val = content[i + 1:j - 1]
            val = re.sub(r"\\nobreakspace\s*(\{\})?", " ", val)
            val = re.sub(r"\s+", " ", val).strip()
            out[k] = val
    return out


def arg(s, cmd):
    m = re.search(r"\\" + cmd + r"(?:\[[^\]]*\])?\{", s)
    i, depth = m.end(), 1
    j = i
    while depth:
        depth += {"{": 1, "}": -1}.get(s[j], 0)
        j += 1
    return s[i:j - 1]


def prepare():
    s = open(TEX).read()
    lab = labels()
    title = arg(s, "title")
    authors = [a for a in re.findall(r"\\author\[[^\]]*\]\{([^}\\]+)", s)]
    email = re.search(r"\\ead\{([^}]+)\}", s).group(1)
    org = re.search(r"organization=\{([^}]+)\}", s).group(1)
    city = re.search(r"city=\{([^}]+)\}", s).group(1)
    country = re.search(r"country=\{([^}]+)\}", s).group(1)
    abstract = s[s.index("\\begin{abstract}") + 16:s.index("\\end{abstract}")].strip()
    keywords = s[s.index("\\begin{keyword}") + 15:s.index("\\end{keyword}")].strip()
    keywords = "; ".join(k.strip() for k in keywords.split("\\sep"))
    body = s[s.index("\\section{Introduction}"):s.index("\\bibliographystyle")]

    # numbered captions and resolved cross-references
    def caption(env):
        def fix(m):
            block = m.group(0)
            lm = re.search(r"\\label\{([^}]+)\}", block)
            if not lm:
                return block
            n = lab.get(lm.group(1), "?")
            word = "Table" if env == "table" else "Figure"
            return block.replace("\\caption{", f"\\caption{{{word} {n}. ", 1)
        return fix
    for env in ("table", "figure"):
        body = re.sub(r"\\begin\{" + env + r"\}.*?\\end\{" + env + r"\}", caption(env), body, flags=re.S)
    body = re.sub(r"\\ref\{([^}]+)\}", lambda m: lab.get(m.group(1), "??"), body)
    body = body.replace("Fig.~", "Fig. ").replace("Table~", "Table ").replace("Section~", "Section ").replace("Sections~", "Sections ")

    # constructs pandoc does not read well
    def fix_tab_spec(spec):
        s = re.sub(r"@\{\\extracolsep\{[^}]*\}\}", "", spec)
        s = re.sub(r"@\{[^}]*\}", "", s)
        s = re.sub(r"p\{[^}]*\}", "l", s)
        s = s.replace("Y", "l").strip()
        return s

    for env in ("tabularx", "tabular\\*"):
        pattern = r"\\begin\{" + env + r"\}\{\\textwidth\}\{(.*?)\}\s*$"
        body = re.sub(pattern, lambda m: "\\begin{tabular}{" + fix_tab_spec(m.group(1)) + "}", body, flags=re.MULTILINE)
    body = body.replace("\\end{tabularx}", "\\end{tabular}")
    body = body.replace("\\end{tabular*}", "\\end{tabular}")
    body = re.sub(r"\\cmidrule(\([^)]*\))?\{[^}]*\}", "", body)
    body = re.sub(r"\\setlength\{\\tabcolsep\}\{[^}]*\}", "", body)
    body = re.sub(r"\\paragraph\{([^}]*)\}\s*", lambda m: "\\textbf{" + m.group(1).rstrip(".") + ".} ", body)
    body = re.sub(r"\\begin\{enumerate\}(.*?)\\end\{enumerate\}",        # labelled list (Q1-Q4) as plain paragraphs
                  lambda m: re.sub(r"\\item\[([^\]]+)\]\s*", r"\n\n\\textbf{\1} ", m.group(1)), body, flags=re.S)
    body = body.replace("\\FloatBarrier", "")

    front = (f"\\begin{{center}}\n{', '.join(authors[:-1])} and {authors[-1]}\n\n"
             f"\\emph{{{org}, {city}, {country}}}\n\nCorresponding author: {authors[0]} ({email})\n\\end{{center}}\n\n"
             f"\\section*{{Abstract}}\n{abstract}\n\n\\noindent\\textbf{{Keywords:}} {keywords}\n\n")
    tail = "\n"                                          # pandoc adds the reference list and its heading
    open(TMP, "w").write("\\documentclass{article}\n\\usepackage{graphicx}\n\\title{" + title + "}\n\\begin{document}\n\\maketitle\n"
                         + front + body + tail + "\\end{document}\n")


def convert():
    subprocess.run(["pandoc", TMP, "-f", "latex", "-t", "docx", "-o", OUT, "--number-sections", "--citeproc",
                    "--bibliography", os.path.join(HERE, "latex", "refs.bib"),
                    "--csl", os.path.join(HERE, "elsevier-with-titles.csl"),
                    "--resource-path", os.path.join(HERE, "figures") + ":" + os.path.join(HERE, "latex"),
                    "-M", "reference-section-title=References", "-M", "link-citations=true"], check=True)


def tidy():
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    d = Document(OUT)
    for name in ("Normal", "Body Text", "First Paragraph", "Compact", "Bibliography", "Abstract"):
        if name in [st.name for st in d.styles]:
            st = d.styles[name]
            st.font.name, st.font.size = "Times New Roman", Pt(12)
            st.paragraph_format.line_spacing = 1.5
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        if name in [st.name for st in d.styles]:
            d.styles[name].font.name = "Times New Roman"
    for t in d.tables:                                   # visible grid, smaller text
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{edge}")
            e.set(qn("w:val"), "single"), e.set(qn("w:sz"), "4"), e.set(qn("w:color"), "808080")
            borders.append(e)
        t._tbl.tblPr.append(borders)
        width = t._tbl.tblPr.find(qn("w:tblW"))
        if width is None:
            width = OxmlElement("w:tblW")
            t._tbl.tblPr.append(width)
        width.set(qn("w:type"), "pct"), width.set(qn("w:w"), "5000")      # full text width
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing = 1.0
                    for r in p.runs:
                        r.font.size = Pt(9.5)
    for p in d.paragraphs:                               # centre figures
        if p._p.xpath(".//w:drawing"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.save(OUT)


if __name__ == "__main__":
    prepare()
    convert()
    tidy()
    os.remove(TMP)
    print("wrote", OUT, file=sys.stderr)
